import csv
import json
import os
import re
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Literal
from zipfile import ZipFile
from xml.etree import ElementTree

from google import genai
from google.genai import errors
from google.genai import types
from openpyxl import load_workbook
from pypdf import PdfReader

from app.schemas.playbook import RuleTemplate
from app.schemas.source import SourceDocument
from app.services.vault_service import VaultService

ExtractionMode = Literal["hybrid", "llm", "heuristic"]

WORDPROCESSINGML_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
SUPPORTED_SUFFIXES = {".xlsx", ".csv", ".docx", ".pdf"}


@dataclass(frozen=True)
class ParsedDocument:
    path: Path
    text: str
    tables: list[list[list[str]]]


class PlaybookExtractionError(RuntimeError):
    pass


class PlaybookExtractionService:
    def __init__(
        self,
        playbook_id: str,
        mode: ExtractionMode = "hybrid",
        gemini_api_key: str | None = None,
        google_cloud_project: str | None = None,
        google_cloud_location: str = "europe-west4",
        gemini_model: str = "gemini-2.5-flash",
        temperature: float = 0.4,
    ) -> None:
        self.playbook_id = playbook_id
        self.mode = mode
        self.gemini_api_key = gemini_api_key or os.getenv("GEMINI_API_KEY")
        self.google_cloud_project = google_cloud_project or os.getenv("GOOGLE_CLOUD_PROJECT")
        self.google_cloud_location = google_cloud_location or os.getenv("GOOGLE_CLOUD_LOCATION", "europe-west4")
        self.gemini_model = gemini_model
        self.temperature = temperature

    def extract_rules(self, source_paths: list[Path]) -> list[RuleTemplate]:
        documents = [parse_document(path) for path in source_paths]
        if self.mode in ("hybrid", "llm") and self._can_use_gemini():
            rules = self._extract_with_gemini(documents)
            if rules:
                return self._dedupe_rule_ids(rules)
        if self.mode == "llm":
            raise PlaybookExtractionError(
                "LLM extraction requested, but Gemini is not configured. "
                "Set GOOGLE_CLOUD_PROJECT for Vertex AI ADC, or GEMINI_API_KEY for API-key fallback."
            )
        return self._dedupe_rule_ids(self._extract_with_heuristics(documents))

    def _extract_with_heuristics(self, documents: list[ParsedDocument]) -> list[RuleTemplate]:
        rules: list[RuleTemplate] = []
        for document in documents:
            rules.extend(extract_rules_from_tables(document, self.playbook_id))
            rules.extend(extract_rules_from_text(document, self.playbook_id))
        return rules

    def _extract_with_gemini(self, documents: list[ParsedDocument]) -> list[RuleTemplate]:
        prompt = build_gemini_prompt(documents, self.playbook_id)
        client = self._gemini_client()
        response = None
        for attempt in range(4):
            try:
                response = client.models.generate_content(
                    model=self.gemini_model,
                    contents=prompt,
                    config=types.GenerateContentConfig(
                        temperature=self.temperature,
                        response_mime_type="application/json",
                    ),
                )
                break
            except errors.APIError as exc:
                if getattr(exc, "code", None) != 429 or attempt == 3:
                    raise
                time.sleep(2**attempt)
        if response is None:
            raise PlaybookExtractionError("Gemini extraction failed without a response.")
        text = response.text
        if not text:
            raise PlaybookExtractionError("Gemini returned an empty extraction response.")
        decoded = json.loads(text)
        raw_rules = decoded["rules"] if isinstance(decoded, dict) and "rules" in decoded else decoded
        if not isinstance(raw_rules, list):
            raise PlaybookExtractionError("LLM response did not contain a rule list.")
        return [RuleTemplate.model_validate(rule) for rule in raw_rules]

    def _can_use_gemini(self) -> bool:
        return bool(self.google_cloud_project or self.gemini_api_key)

    def _gemini_client(self) -> genai.Client:
        if self.google_cloud_project:
            return genai.Client(
                vertexai=True,
                project=self.google_cloud_project,
                location=self.google_cloud_location,
            )
        if self.gemini_api_key:
            return genai.Client(api_key=self.gemini_api_key)
        raise PlaybookExtractionError("Gemini is not configured.")

    @staticmethod
    def _dedupe_rule_ids(rules: list[RuleTemplate]) -> list[RuleTemplate]:
        seen: dict[str, int] = {}
        deduped: list[RuleTemplate] = []
        for rule in rules:
            base_id = rule.rule_id or VaultService.slugify_topic(rule.topic)
            count = seen.get(base_id, 0)
            seen[base_id] = count + 1
            rule.rule_id = base_id if count == 0 else f"{base_id}-{count + 1}"
            deduped.append(rule)
        return deduped


def parse_document(path: Path) -> ParsedDocument:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise PlaybookExtractionError(f"Unsupported playbook source format: {path}")
    if suffix == ".xlsx":
        return parse_xlsx(path)
    if suffix == ".csv":
        return parse_csv(path)
    if suffix == ".docx":
        return parse_docx(path)
    if suffix == ".pdf":
        return parse_pdf(path)
    raise PlaybookExtractionError(f"Unsupported playbook source format: {path}")


def parse_xlsx(path: Path) -> ParsedDocument:
    workbook = load_workbook(path, read_only=True, data_only=True)
    tables: list[list[list[str]]] = []
    text_parts: list[str] = []
    for sheet in workbook.worksheets:
        table = [[normalize_cell(value) for value in row] for row in sheet.iter_rows(values_only=True)]
        table = [row for row in table if any(row)]
        if table:
            tables.append(table)
            text_parts.append(f"Sheet: {sheet.title}")
            text_parts.extend(" | ".join(row) for row in table)
    return ParsedDocument(path=path, text="\n".join(text_parts), tables=tables)


def parse_csv(path: Path) -> ParsedDocument:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        table = [[cell.strip() for cell in row] for row in csv.reader(handle)]
    table = [row for row in table if any(row)]
    return ParsedDocument(path=path, text="\n".join(" | ".join(row) for row in table), tables=[table] if table else [])


def parse_docx(path: Path) -> ParsedDocument:
    with ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read("word/document.xml"))
    paragraphs = []
    for paragraph in root.iter(WORDPROCESSINGML_NS + "p"):
        text = "".join(node.text or "" for node in paragraph.iter(WORDPROCESSINGML_NS + "t")).strip()
        if text:
            paragraphs.append(text)
    return ParsedDocument(path=path, text="\n".join(paragraphs), tables=[])


def parse_pdf(path: Path) -> ParsedDocument:
    reader = PdfReader(path)
    pages = [page.extract_text() or "" for page in reader.pages]
    return ParsedDocument(path=path, text="\n\n".join(page.strip() for page in pages if page.strip()), tables=[])


def extract_rules_from_tables(document: ParsedDocument, playbook_id: str) -> list[RuleTemplate]:
    rules: list[RuleTemplate] = []
    for table in document.tables:
        if len(table) < 2:
            continue
        header_index, aliases = detect_header(table)
        if header_index is None:
            continue
        headers = table[header_index]
        for row_number, row in enumerate(table[header_index + 1 :], start=header_index + 2):
            values = {aliases[index]: row[index] for index in range(min(len(headers), len(row))) if index in aliases}
            topic = values.get("topic") or values.get("clause_name")
            if not topic:
                continue
            fallback_positions = [values[key] for key in sorted(values) if key.startswith("fallback") and values[key]]
            source = SourceDocument(filename=display_path(document.path), location=f"Row {row_number}")
            rules.append(
                RuleTemplate(
                    playbook_id=playbook_id,
                    rule_id=VaultService.slugify_topic(topic),
                    topic=topic,
                    standard_position=values.get("standard_position") or None,
                    fallback_positions=fallback_positions,
                    red_line=values.get("red_line") or None,
                    decision_logic=values.get("decision_logic") or None,
                    escalation_logic=values.get("escalation_logic") or None,
                    rationale=values.get("rationale") or None,
                    negotiation_tips=[values["watch_for"]] if values.get("watch_for") else [],
                    suggested_language=values.get("suggested_language") or None,
                    status="approved",
                    source_documents=[source],
                )
            )
    return rules


def detect_header(table: list[list[str]]) -> tuple[int | None, dict[int, str]]:
    best_index: int | None = None
    best_aliases: dict[int, str] = {}
    for row_index, row in enumerate(table[:10]):
        aliases = {}
        for column_index, value in enumerate(row):
            alias = header_alias(value)
            if alias:
                aliases[column_index] = alias
        if len(aliases) > len(best_aliases):
            best_index = row_index
            best_aliases = aliases
    if "topic" not in set(best_aliases.values()) and "clause_name" not in set(best_aliases.values()):
        return None, {}
    return best_index, best_aliases


def header_alias(value: str) -> str | None:
    normalized = re.sub(r"[^a-z0-9]+", " ", value.lower()).strip()
    if normalized in {"clause name", "topic", "issue", "rule", "clause"}:
        return "topic"
    if normalized in {"preferred position", "standard position", "default position", "position"}:
        return "standard_position"
    if normalized.startswith("fallback"):
        suffix = re.sub(r"[^0-9]+", "", normalized) or "1"
        return f"fallback_{suffix}"
    if normalized in {"red line", "redline", "must not accept"}:
        return "red_line"
    if normalized in {"escalation trigger", "escalation logic", "when to escalate"}:
        return "escalation_logic"
    if normalized in {"decision logic", "conditions", "if then"}:
        return "decision_logic"
    if normalized in {"why it matters summary", "why it matters", "rationale", "explanation"}:
        return "rationale"
    if normalized in {"what to watch for", "watch out", "negotiation tips", "tips"}:
        return "watch_for"
    if normalized in {"suggested language", "sample language", "clause language"}:
        return "suggested_language"
    if normalized in {"clause", "clause number", "clause #"}:
        return "clause_number"
    return None


def extract_rules_from_text(document: ParsedDocument, playbook_id: str) -> list[RuleTemplate]:
    sections = split_clause_sections(document.text)
    if not sections:
        return []
    rules = []
    for index, (title, body) in enumerate(sections, start=1):
        parsed = parse_labeled_rule_body(body)
        topic = title.strip() or f"Rule {index}"
        rules.append(
            RuleTemplate(
                playbook_id=playbook_id,
                rule_id=VaultService.slugify_topic(topic),
                topic=topic,
                standard_position=parsed.get("standard_position"),
                fallback_positions=parsed.get("fallback_positions", []),
                red_line=parsed.get("red_line"),
                decision_logic=parsed.get("decision_logic"),
                escalation_logic=parsed.get("escalation_logic"),
                rationale=parsed.get("rationale"),
                negotiation_tips=parsed.get("negotiation_tips", []),
                suggested_language=parsed.get("suggested_language"),
                status="approved",
                source_documents=[SourceDocument(filename=display_path(document.path), location=f"Clause section {index}")],
            )
        )
    return rules


def split_clause_sections(text: str) -> list[tuple[str, str]]:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    starts: list[tuple[int, str]] = []
    for index, line in enumerate(lines):
        match = re.match(r"^(?:Clause\s+)?(\d+)[\).:\s-]+(.+)$", line, flags=re.IGNORECASE)
        if match and len(match.group(2)) < 160:
            starts.append((index, match.group(2).strip()))
            continue
        match = re.match(r"^Clause\s+\d+:\s*(.+)$", line, flags=re.IGNORECASE)
        if match:
            starts.append((index, match.group(1).strip()))
    sections = []
    for position, (line_index, title) in enumerate(starts):
        next_index = starts[position + 1][0] if position + 1 < len(starts) else len(lines)
        body = "\n".join(lines[line_index + 1 : next_index])
        if body:
            sections.append((title, body))
    return sections


def parse_labeled_rule_body(body: str) -> dict:
    lines = [line.strip() for line in body.splitlines() if line.strip()]
    result = {"fallback_positions": [], "negotiation_tips": []}
    current_label: str | None = None
    for line in lines:
        label, value = split_label(line)
        if label:
            current_label = label
            if value:
                add_labeled_value(result, label, value)
            continue
        if current_label:
            add_labeled_value(result, current_label, line)
    if not any(result.values()):
        result["standard_position"] = body.strip()
    return result


def split_label(line: str) -> tuple[str | None, str]:
    normalized = line.lower().replace("🔴", "").strip()
    labels = {
        "why it matters:": "rationale",
        "what to watch for:": "watch_for",
        "preferred": "standard_marker",
        "fallback 1": "fallback_marker",
        "fallback 2": "fallback_marker",
        "red line": "red_line_marker",
        "escalation trigger:": "escalation_logic",
        "suggested language:": "suggested_language",
        "decision logic:": "decision_logic",
    }
    for prefix, label in labels.items():
        if normalized.startswith(prefix):
            return label, line[len(prefix) :].strip(" :")
    return None, ""


def add_labeled_value(result: dict, label: str, value: str) -> None:
    if label == "standard_marker":
        result["standard_position"] = append_text(result.get("standard_position"), value)
    elif label == "fallback_marker":
        if value:
            result.setdefault("fallback_positions", []).append(value)
    elif label == "red_line_marker":
        result["red_line"] = append_text(result.get("red_line"), value)
    elif label == "watch_for":
        result.setdefault("negotiation_tips", []).append(f"What to watch for: {value}")
    elif label in {"rationale", "escalation_logic", "suggested_language", "decision_logic"}:
        result[label] = append_text(result.get(label), value)


def append_text(existing: str | None, value: str) -> str:
    if not existing:
        return value
    return f"{existing} {value}"


def build_gemini_prompt(documents: list[ParsedDocument], playbook_id: str) -> str:
    document_text = "\n\n".join(
        f"SOURCE: {display_path(document.path)}\n{document.text[:25000]}" for document in documents
    )
    return f"""
Extract legal playbook rules from the provided source material.

Return only JSON with this shape:
{{
  "rules": [
    {{
      "playbook_id": "{playbook_id}",
      "rule_id": "stable-slug",
      "topic": "Rule topic",
      "standard_position": "Preferred or standard position, if available",
      "fallback_positions": ["Fallback positions"],
      "red_line": "Unacceptable position, if available",
      "decision_logic": "Decision logic, if available",
      "escalation_logic": "Escalation trigger, if available",
      "rationale": "Plain-language explanation, if available",
      "negotiation_tips": ["Practical watch-outs or negotiation tips"],
      "suggested_language": "Suggested clause language, if available",
      "status": "approved",
      "source_documents": [
        {{"filename": "source path", "location": "where in the source this rule came from"}}
      ]
    }}
  ]
}}

Do not assume a fixed number of rules. Infer the rules from headings, tables, labels,
and surrounding context. Preserve auditability by filling source_documents.

{document_text}
""".strip()


def normalize_cell(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def display_path(path: Path) -> str:
    try:
        return str(path.resolve().relative_to(Path.cwd()))
    except ValueError:
        return str(path)
